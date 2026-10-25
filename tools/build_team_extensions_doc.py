"""Build the member extension handbook from the reviewed Markdown direction cards.

Optional publishing helper; does not participate in mod compilation.
Requires python-docx. Render the output and review every page before distribution.
"""
import argparse
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
REPOSITORY='https://github.com/Twilight429/mc-companion-agent'
SOURCE=REPOSITORY+'/blob/1.21.1/docs/'

def face(style, east='宋体', latin='Calibri', size=11.5):
    style.font.name=latin;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
    for key,value in [('eastAsia',east),('ascii',latin),('hAnsi',latin)]:fonts.set(qn('w:'+key),value)

def inline(p,text):
    for fragment in re.split(r'(\*\*.*?\*\*|`[^`]+`)',text):
        if not fragment:continue
        if fragment.startswith('**'):
            p.add_run(fragment[2:-2]).bold=True
        elif fragment.startswith('`'):
            r=p.add_run(fragment[1:-1]);r.font.name='Consolas';r.font.size=Pt(10)
        else:p.add_run(fragment)

def para(doc,text='',bold=False):
    p=doc.add_paragraph();inline(p,text)
    if bold:
        for r in p.runs:r.bold=True
    return p

def link(doc,label,url):
    p=doc.add_paragraph();element=OxmlElement('w:hyperlink')
    element.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'20547A');pr.append(color)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);element.append(r);p._p.append(element)
    return p

def table(doc,headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    props.append(borders)
    for idx,values in enumerate([headers]+rows):
        row=t.rows[0] if idx==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if idx==0:trpr.append(OxmlElement('w:tblHeader'))
        for col,(cell,value,w) in enumerate(zip(row.cells,values,widths)):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for side,val in [('top',85),('bottom',85),('left',95),('right',95)]:
                item=OxmlElement('w:'+side);item.set(qn('w:w'),str(val));item.set(qn('w:type'),'dxa');margins.append(item)
            tcpr.append(margins)
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'294253' if idx==0 else ('F0F4F7' if idx%2==0 else 'FFFFFF'));tcpr.append(shd)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.12
            if col==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(value);r.font.size=Pt(10.5);r.bold=(idx==0)
            if idx==0:r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def fraction(doc):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    mathpara=OxmlElement('m:oMathPara');math=OxmlElement('m:oMath')
    def run(text):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
    math.append(run('p = '));f=OxmlElement('m:f');num=OxmlElement('m:num');den=OxmlElement('m:den')
    num.append(run('a + h'));den.append(run('a + b + h + u'));f.append(num);f.append(den);math.append(f)
    mathpara.append(math);p._p.append(mathpara)

def section(doc,title):
    doc.add_page_break();doc.add_heading(title,level=1)

def build(output):
    text=(ROOT/'docs/TEAM_EXTENSIONS.md').read_text(encoding='utf-8')
    source={}
    for letter,name,body in re.findall(r'^## ([A-E])\. ([^\n]+)\n(.*?)(?=^## |\Z)',text,re.M|re.S):source[letter]=(name,body.strip())
    if set(source)!=set('ABCDE'):raise ValueError('Expected all five reviewed extension directions')
    doc=Document();s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=Cm(1.7);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
    normal=doc.styles['Normal'];face(normal);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(5)
    for name,size in [('Title',23),('Subtitle',12),('Heading 1',16),('Heading 2',12)]:
        style=doc.styles[name];face(style,'微软雅黑','Calibri',size);style.font.bold=name!='Subtitle';style.font.italic=False
        style.paragraph_format.space_before=Pt(7);style.paragraph_format.space_after=Pt(7)
        style.paragraph_format.keep_with_next=True
    for name in ('Title','Subtitle','Heading 1','Heading 2'):
        pr=doc.styles[name].element.get_or_add_pPr()
        for item in list(pr):
            if item.tag==qn('w:pBdr'):pr.remove(item)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for val in ('第 ',None,' 页'):
        if val is None:
            f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
        else:footer.add_run(val).font.size=Pt(9)
    doc.core_properties.title='MC陪玩Agent组员扩展档案';doc.core_properties.author='项目组';doc.core_properties.subject='扩展方向认领与实践验收'
    doc.add_paragraph('MC陪玩Agent组员扩展档案',style='Title')
    doc.add_paragraph('方向认领与实践验收',style='Subtitle')
    para(doc,'2026年10月25日  |  面向四人项目组  |  源码基线 85e0b24')
    para(doc,'我们要让一个真实存在于 Minecraft 世界中的 NPC 理解玩家意图，与玩家共同完成目标，并利用可追溯的反馈改善协作。每位组员从本档案中认领一个主方向，先做出可演示、可复现的小功能，再深化方法。')
    para(doc,'当前载体是 Minecraft Java 1.21.1 与 Fabric，沿用 Numen 的唯一 Agent 大脑、游戏感知、身体与权限。已有建筑投影和协作、确认式玩家偏好、生存与交互接口；本轮扩展不要求先训练RL或接入多模态大模型。')
    doc.add_heading('方向选择总览',level=2)
    table(doc,['编号','扩展方向','第一份可见成果','建议'],[
        ['A','玩家偏好与自适应协作','记住确认的偏好并按任务调整帮助','负责人或1人'],
        ['B','生存补给与采集探索','完成一种真实的物资补给闭环','1人'],
        ['C','执行经验与失败恢复','遇到熟悉错误时少走重复弯路','1人'],
        ['D','陪玩互动与共同目标','能适时建议配合并响应拒绝','1人'],
        ['E','实验评测与状态观测','回看规划到实际结果的完整时间线','可替换B C D之一'],
    ],[1.3,5.2,7.7,3.3])
    para(doc,'建议负责人维护集成并推进A，其余三位从B至E各选一个。五个方向是自选菜单，不要求四人全部完成；每位实现者都要验证自己的功能。')
    para(doc,'当前团队快照已通过2037项单元测试和24项Fabric端到端测试。端到端使用模拟模型驱动真实服务器，不能据此认定真实LLM的长期陪玩效果已被验证。')
    para(doc,'下文数值是拟定验收目标，不是已取得的结果。先做预实验，再冻结场景、预算和成功定义；失败也必须进入报告。',bold=True)

    for letter in 'ABCDE':
        name,body=source[letter];section(doc,'方向 '+letter+' '+name)
        for chunk in body.split('\n\n'):
            chunk=chunk.strip().replace('\n',' ')
            if letter=='E':chunk=chunk.replace('成功率 `成功次数/总尝试数`','成功率按成功次数除以总尝试数计算')
            if not chunk:continue
            if letter=='A' and chunk.startswith('**可做的方法。**'):
                lead,tail=chunk.split('`p=(a+h)/(a+b+h+u)`',1)
                para(doc,lead.rstrip('，'));fraction(doc);para(doc,tail.lstrip('；'))
            else:para(doc,chunk)

    section(doc,'共同交付与推进安排')
    para(doc,'每项扩展至少交付代码、一个可复现场景、对照结果、一段短演示以及失败复盘。不能只提交提示词或AI生成的说明，也不能只凭NPC说“完成了”验收。')
    doc.add_heading('分阶段推进',level=2)
    table(doc,['阶段','成员交付','负责人验收'],[
        ['第1周','跑通构建并认领方向 复现一个缺陷 提交一页接口和实验设计','固定基线与共同记录字段 审查接口边界'],
        ['第2周','完成最小功能 一个场景和相应正确性测试 录制短演示','逐个接入 排查第二套大脑或绕过权限的实现'],
        ['第3至4周','完成对照场景 保留失败数据 根据结果改进一轮','固定评测条件 安排另一位成员交叉复现'],
        ['后续阶段','依据结果选择一个点深化 避免持续堆加功能','集成可展示主线 保留课程与实践缓冲'],
    ],[2.3,8.1,7.1])
    para(doc,'这是认领后的相对节奏，不重置课程十周截止时间。按剩余周数压缩，至少保留2至3周用于集成与课程实践。方向卡中的1至2周仅指首次闭环估计。')
    doc.add_heading('统一实验记录',level=2)
    para(doc,'每次运行记录run_id、代码提交、模型标识、场景种子、起点、库存、权限、时间与调用预算、终止原因。任务成功由实际库存或世界状态核验。超时、死亡、人工干预与重试都要保留，不能覆盖旧失败。')
    para(doc,'明确分开三类证据：模拟模型用于检验接口与控制流程；真实模型实验用于评估规划与行动；自愿玩家体验用于探索配合感和烦扰感。相同模型和起点下做配对对照，小样本只作探索性结论。')
    doc.add_heading('合并检查',level=2)
    para(doc,'每个PR回答：玩家现在能做什么；复用了哪些原生接口；世界状态如何验收；相对哪个基线；失败了哪些情况；AI生成部分是否经本人审阅。网络包、持久化字段、权限和核心任务变更先找负责人评审。')

    section(doc,'认领记录与开始实践')
    para(doc,'姓名与具体认领由组内讨论后填写。每人优先选择一个主方向；跨方向依赖只约定最小数据契约，不同时重写他人的模块。')
    table(doc,['成员','主方向','本次最小目标','功能分支'],[
        ['负责人','','',''],['成员二','','',''],['成员三','','',''],['成员四','','',''],
    ],[2.1,3.1,8.5,3.8])
    doc.add_heading('开始前完成三件事',level=2)
    para(doc,'第一，获得仓库访问权限，克隆1.21.1分支，安装JDK21，按构建指南生成资源并编译Fabric模组。在自己的隔离游戏实例中配置模型，不依赖负责人电脑的Key、存档或路径。')
    para(doc,'第二，阅读根目录AGENTS.md、CONTRIBUTING.md与所认领模块，先检查现有公开API和技能；对照自己要做的功能，写一个玩家能直接观察到差异的例子。')
    para(doc,'第三，从稳定基线开个人功能分支，小步实现、运行相应检查、发PR。不要强推，不直接覆盖稳定分支，也不要删除测试来获得通过。')
    doc.add_heading('给AI coding工具的任务描述',level=2)
    para(doc,'我认领方向X，本次只做该方向的最小闭环。请先读取团队指南与AGENTS.md，检查已有接口，列出改动文件和一个可复现实验，再实现并运行必要检查。复用唯一Numen大脑、原生任务与权限；不得修改私人存档、伪造资源或把Key写入代码。最后列出真实执行证据、失败情况和剩余限制。')
    doc.add_heading('项目入口与来源',level=2)
    link(doc,'团队源码与项目入口',REPOSITORY)
    link(doc,'Gitee仓库入口','https://gitee.com/X1nze/mc-companion-agent')
    link(doc,'构建安装与协作指南',SOURCE+'TEAM_BUILD.md')
    link(doc,'扩展方向原始说明',SOURCE+'TEAM_EXTENSIONS.md')
    link(doc,'团队快照验证记录',SOURCE+'TEAM_VERIFICATION.md')
    para(doc,'上游工程为Dwinovo与贡献者的Numen，代码基线daae0f1，沿用LGPL 3.0许可证并保留署名。团队快照使用独立占位图标，不包含个人Key、存档、私人YSM模型或用户下载的建筑作品。完整来源与素材边界见仓库TEAM_PROVENANCE.md。')
    link(doc,'Numen上游工程','https://github.com/Dwinovo/minecraft-numen')
    output.parent.mkdir(parents=True,exist_ok=True);doc.save(output)
    print(str(output.resolve()))

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--output',type=Path,default=ROOT/'docs/handbook/MC陪玩Agent_组员扩展档案_20261025.docx')
    build(parser.parse_args().output)
